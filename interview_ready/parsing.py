"""Extract plain text from uploaded CV/JD files. Nothing is logged."""
import io
from pathlib import Path


def extract_text(name: str, data: bytes) -> str:
    suffix = Path(name).suffix.lower()
    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        return "\n".join((page.extract_text() or "") for page in reader.pages)
    if suffix == ".docx":
        from docx import Document

        doc = Document(io.BytesIO(data))
        lines = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                lines.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(lines)
    if suffix in (".txt", ".md", ""):
        return data.decode("utf-8", errors="replace")
    raise ValueError(f"Unsupported file type: {suffix}")
